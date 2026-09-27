"""Optional fixture regeneration: uv run --with python-pptx python test/generate-preview-slide.py."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

presentation = Presentation()
for title in ("Cloudex PPT Preview", "Second slide"):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    box.text_frame.text = title
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
target = Path(__file__).parent / "fixtures" / "preview.pptx"
target.parent.mkdir(exist_ok=True)
presentation.save(target)
